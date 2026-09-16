# -*- coding: utf-8 -*-
"""検証用：既知のテストデータを流し込んだコピーを作る"""
import shutil, datetime
from openpyxl import load_workbook

shutil.copy("magnet-nail-mri_workbook.xlsx", "verify.xlsx")
wb = load_workbook("verify.xlsx")

# --- 01 試料マスタ：3件 ---
m = wb["01_試料マスタ"]
samples = [("S-01","製品A","マグネットネイル"), ("S-02","製品B","カラージェル"), ("S-03","製品C","マグネットネイル")]
for i,(sid,pn,cat) in enumerate(samples):
    r = 2+i
    m.cell(row=r, column=1, value=sid); m.cell(row=r, column=2, value=pn); m.cell(row=r, column=3, value=cat)

# --- 02 偏向角：S-01 3T に 3件、S-01 1.5T に 2件、S-03 3T に 1件(46度=要確認) ---
d = wb["02_偏向角_記録"]
defl = [("D-1","3T","S-01",0.085,18.0),("D-2","3T","S-01",0.085,20.0),("D-3","3T","S-01",0.085,22.0),
        ("D-4","1.5T","S-01",0.085,11.0),("D-5","1.5T","S-01",0.085,13.0),
        ("D-6","3T","S-03",0.090,46.0)]
for i,(mid,dev,sid,mass,th) in enumerate(defl):
    r = 2+i
    d.cell(row=r,column=1,value=mid); d.cell(row=r,column=3,value=dev); d.cell(row=r,column=4,value=sid)
    d.cell(row=r,column=5,value=i+1); d.cell(row=r,column=6,value=mass); d.cell(row=r,column=7,value=th)

# --- 04 アーチファクト：GRE/SE, 距離, TE ---
a = wb["04_アーチファクト_記録"]
art = [  # (装置, seq, TE, 距離, 角度, 面積)
    ("3T","GRE",15,0,0,120.0), ("3T","GRE",15,0,0,130.0), ("3T","GRE",15,2,0,60.0),
    ("3T","GRE",15,5,0,8.0),   ("3T","SE",15,0,0,30.0),   ("1.5T","GRE",15,0,0,70.0),
    ("1.5T","GRE",15,0,0,80.0),("1.5T","GRE",15,5,0,4.0), ("3T","GRE",30,0,0,230.0),
]
for i,(dev,seq,te,dis,ang,area) in enumerate(art):
    r = 2+i
    a.cell(row=r,column=1,value=f"A-{i+1}"); a.cell(row=r,column=3,value=dev); a.cell(row=r,column=4,value=seq)
    a.cell(row=r,column=5,value=te); a.cell(row=r,column=8,value=dis); a.cell(row=r,column=9,value=ang)
    a.cell(row=r,column=10,value="S-01"); a.cell(row=r,column=11,value=1); a.cell(row=r,column=12,value=area)

# --- 06 発熱：40℃未満 と 50〜55℃ ---
h = wb["06_発熱_記録"]
h.cell(row=2,column=1,value="H-1"); h.cell(row=2,column=7,value="S-01")
for c_ in range(8,13): h.cell(row=2,column=c_,value="未変色")
h.cell(row=3,column=1,value="H-2"); h.cell(row=3,column=7,value="S-03")
for c_,v in zip(range(8,13), ["変色","変色","変色","未変色","未変色"]): h.cell(row=3,column=c_,value=v)

# --- 07 延期症例：延期3件(頭部2,手1)、非延期2件、未受検1件 ---
c = wb["07_延期症例"]
base = datetime.date(2026,1,1)
cases = [  # (依頼日offset, 実施日offset, 部位, 延期, 未受検)
    (0, 30, "頭部", "有", "非該当"),
    (0, 20, "頭部", "有", "非該当"),
    (0, 40, "手・手関節", "有", "該当"),
    (0,  8, "体幹部", "無", "非該当"),
    (0, 10, "頭部", "無", "非該当"),
]
for i,(o1,o2,part,post,drop) in enumerate(cases):
    r = 2+i
    c.cell(row=r,column=1,value=f"C-{i+1}")
    c.cell(row=r,column=2,value=base+datetime.timedelta(days=o1))
    c.cell(row=r,column=4,value=base+datetime.timedelta(days=o2))
    c.cell(row=r,column=5,value=part); c.cell(row=r,column=6,value="マグネットネイル")
    c.cell(row=r,column=7,value=post); c.cell(row=r,column=9,value=drop)

# --- 09 スクリーニング：TP3 FP1 FN1 TN3 ---
s = wb["09_スクリーニング検証"]
scr = [("陽性","磁性あり")]*3 + [("陽性","磁性なし")]*1 + [("陰性","磁性あり")]*1 + [("陰性","磁性なし")]*3
for i,(res,gt) in enumerate(scr):
    r = 5+i
    s.cell(row=r,column=1,value=f"S-{i+1:02d}"); s.cell(row=r,column=2,value=res); s.cell(row=r,column=3,value=gt)

wb.save("verify.xlsx")
print("verify.xlsx created")
