# -*- coding: utf-8 -*-
"""マグネットネイル×MRI 実験記録・集計ワークブック"""
import sys, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

FONT = "Meiryo"
HDR_FILL   = PatternFill("solid", fgColor="2E3A59")
SUB_FILL   = PatternFill("solid", fgColor="E6F3F5")
IN_FILL    = PatternFill("solid", fgColor="FFF9E8")   # 入力セル
CALC_FILL  = PatternFill("solid", fgColor="F2F4F8")   # 自動計算
EX_FILL    = PatternFill("solid", fgColor="EAF1F8")   # 記入例
TITLE_FILL = PatternFill("solid", fgColor="0F1626")
WARN_FILL  = PatternFill("solid", fgColor="FBE3E6")
OK_FILL    = PatternFill("solid", fgColor="E4F3E9")

thin = Side(style="thin", color="D4DAE4")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

def f(sz=10, b=False, color="1F2430", it=False):
    return Font(name=FONT, size=sz, bold=b, color=color, italic=it)

# ---------- 行数 ----------
N_MASTER, N_DEFL, N_ART, N_HEAT, N_CASE, N_SCR, N_LOG = 120, 400, 800, 150, 300, 200, 150

wb = Workbook()
wb.remove(wb.active)

def sheet(name):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    return ws

def headers(ws, cols, row=1, fill=HDR_FILL):
    """cols: [(見出し, 幅, 種別)] 種別 in {'in','calc','-'}"""
    for i, (h, w, kind) in enumerate(cols, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = f(10, True, "FFFFFF"); c.fill = fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 32

def paint(ws, cols, r0, r1):
    for i, (_, _, kind) in enumerate(cols, start=1):
        if kind == "-":
            continue
        fill = IN_FILL if kind == "in" else CALC_FILL
        for r in range(r0, r1 + 1):
            c = ws.cell(row=r, column=i)
            c.fill = fill; c.font = f(10); c.border = BORDER
            c.alignment = Alignment(vertical="center")

def example(ws, cols, row, vals):
    for i, v in enumerate(vals, start=1):
        if v is None:
            continue
        c = ws.cell(row=row, column=i, value=v)
        c.fill = EX_FILL; c.font = f(10, False, "3A5A80", it=True)

def dv(ws, rng, items, prompt=None):
    d = DataValidation(type="list", formula1='"' + ",".join(items) + '"', allow_blank=True)
    d.error = "リストから選んでください"; d.errorTitle = "入力エラー"
    if prompt:
        d.prompt = prompt; d.promptTitle = "入力"
    ws.add_data_validation(d); d.add(rng)

def dv_range(ws, rng, src):
    d = DataValidation(type="list", formula1=src, allow_blank=True)
    ws.add_data_validation(d); d.add(rng)

def note(ws, cell, text, color="6B7280", sz=9):
    ws[cell] = text; ws[cell].font = f(sz, False, color)

# =====================================================================
# 00_使い方
# =====================================================================
ws = sheet("00_使い方")
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 26
ws.column_dimensions["C"].width = 92
ws["B2"] = "マグネットネイル × MRI　実験記録・集計ワークブック"
ws["B2"].font = f(16, True, "0F1626"); ws.row_dimensions[2].height = 26
ws["B3"] = "PRESENTATION_PLAN.md の案A〜案F に対応。記録シートに入力すると、集計シートが自動で更新されます。"
ws["B3"].font = f(10, False, "6B7280")

rows = [
    ("セルの色", "", "sec"),
    ("うすい黄色", "入力するセル。ここだけ手で埋めます。", ""),
    ("うすい灰色", "自動計算セル。数式が入っています。上書きしないでください。", ""),
    ("うすい青色", "記入例。1行目のデータ行に入れてあります。上書きするか、行ごと削除して構いません。", ""),
    ("", "", "gap"),
    ("シートの構成", "", "sec"),
    ("01_試料マスタ", "供試するネイル試料の台帳。以降の全シートはここの試料IDを参照します。最初に埋めてください。", ""),
    ("02_偏向角_記録", "案A。ASTM F2052 の偏向角法。1測定1行。磁気力は自動計算されます。", ""),
    ("03_偏向角_集計", "02から自動集計。試料×装置ごとの平均・SD・95%信頼区間・45°判定。入力不要。", ""),
    ("04_アーチファクト_記録", "案B。ファントム撮影。1条件1測定1行。撮影条件をすべて残してください。", ""),
    ("05_アーチファクト_集計", "04から自動集計。シーケンス別・TE別・距離別。入力不要。", ""),
    ("06_発熱_記録", "案C。示温ラベルの変色を記録。最高到達温度の範囲が自動で出ます。", ""),
    ("07_延期症例", "案E-1。延期の実態調査。依頼から実施までの日数が自動計算されます。", ""),
    ("08_延期集計", "07から自動集計。延期率・日数の中央値・部位別内訳・未受検件数。入力不要。", ""),
    ("09_スクリーニング検証", "案E-3。簡易磁石テストの感度・特異度。2×2表と指標が自動で出ます。", ""),
    ("10_撮影ログ", "実施記録。日時・装置・立会者。安全管理の記録としても残します。", ""),
    ("11_安全チェックリスト", "撮影のたびに確認する項目。実施前・実施後。", ""),
    ("", "", "gap"),
    ("使う順番", "", "sec"),
    ("1", "01_試料マスタ に供試試料を登録する（試料IDを先に決める）", ""),
    ("2", "07_延期症例 の後ろ向き集計を進める（実験と並行して着手できる）", ""),
    ("3", "10_撮影ログ と 11_安全チェックリスト を用意して実機作業に入る", ""),
    ("4", "02 / 04 / 06 に測定値を入れる。集計シートは自動で更新される", ""),
    ("5", "09 で簡易テストの性能を評価し、運用提案につなげる", ""),
    ("", "", "gap"),
    ("注意", "", "sec"),
    ("数式について", "自動計算セルの数式は範囲を固定しています。行を挿入する場合は、データ行の末尾より内側に挿入してください。", ""),
    ("行数の上限", "02は400行、04は800行、06は150行、07は300行、09は200行まで数式を入れてあります。足りなければ最終行を下方向にコピーしてください。", ""),
    ("検証について", "この環境では Excel / LibreOffice による再計算ができなかったため、数式は Python の Excel 数式評価エンジンで評価して確認しています。初回はお手元の Excel で開いて、集計値が意図どおりか確認してください。", ""),
]
r = 5
for a, b, kind in rows:
    if kind == "gap":
        r += 1; continue
    if kind == "sec":
        ws.cell(row=r, column=2, value=a).font = f(12, True, "1B9AAA")
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        r += 1; continue
    ca = ws.cell(row=r, column=2, value=a); ca.font = f(10, True); ca.alignment = Alignment(vertical="top")
    cb = ws.cell(row=r, column=3, value=b); cb.font = f(10); cb.alignment = Alignment(vertical="top", wrap_text=True)
    ws.row_dimensions[r].height = 30 if len(b) > 60 else 18
    r += 1
for cell, fill in (("B6", IN_FILL), ("B7", CALC_FILL), ("B8", EX_FILL)):
    ws[cell].fill = fill

# =====================================================================
# 01_試料マスタ
# =====================================================================
ws = sheet("01_試料マスタ")
cols = [("試料ID", 11, "in"), ("製品記号", 12, "in"), ("カテゴリ", 18, "in"), ("色", 12, "in"),
        ("ロット", 14, "in"), ("施術日", 12, "in"), ("磁石配列", 12, "in"),
        ("塗布量 (g)", 11, "in"), ("備考", 40, "in")]
headers(ws, cols); paint(ws, cols, 2, N_MASTER + 1)
example(ws, cols, 2, ["S-01", "製品A", "マグネットネイル", "ネイビー", "L2405A", datetime.date(2026, 1, 15), "縦配列", 0.042, "猫目の発色が強いタイプ"])
dv(ws, f"C3:C{N_MASTER+1}", ["透明ジェル", "カラージェル", "ラメ入り", "マグネットネイル", "マニキュア", "その他"])
dv(ws, f"G3:G{N_MASTER+1}", ["未配列", "縦配列", "横配列", "斜め配列"])
for r_ in range(2, N_MASTER + 2):
    ws.cell(row=r_, column=8).number_format = "0.000"
    ws.cell(row=r_, column=6).number_format = "yyyy/mm/dd"
ws.freeze_panes = "A2"

# =====================================================================
# 02_偏向角_記録  (案A / ASTM F2052)
# =====================================================================
ws = sheet("02_偏向角_記録")
cols = [("測定ID", 11, "in"), ("測定日", 12, "in"), ("装置", 9, "in"), ("試料ID", 11, "in"),
        ("反復No", 8, "in"), ("質量 m (g)", 11, "in"), ("偏向角 θ (°)", 12, "in"),
        ("磁気力 F (mN)", 13, "calc"), ("45°判定", 10, "calc"),
        ("測定者", 12, "in"), ("備考", 32, "in")]
headers(ws, cols); paint(ws, cols, 2, N_DEFL + 1)
example(ws, cols, 2, ["D-0001", datetime.date(2026, 2, 3), "3T", "S-01", 1, 0.085, 18.4, None, None, "山田", "ボア入口 ∂B/∂z 最大位置"])
dv(ws, f"C2:C{N_DEFL+1}", ["1.5T", "3T"])
dv_range(ws, f"D2:D{N_DEFL+1}", f"='01_試料マスタ'!$A$2:$A${N_MASTER+1}")
for r_ in range(2, N_DEFL + 2):
    # F(mN) = m[g] * 9.80665 * tanθ   （m[kg]*g*tanθ を mN に換算すると g 単位がそのまま使える）
    ws.cell(row=r_, column=8, value=f'=IF(OR(F{r_}="",G{r_}=""),"",ROUND(F{r_}*9.80665*TAN(RADIANS(G{r_})),3))').number_format = "0.000"
    ws.cell(row=r_, column=9, value=f'=IF(G{r_}="","",IF(G{r_}<45,"適合","要確認"))')
    ws.cell(row=r_, column=7).number_format = "0.0"
    ws.cell(row=r_, column=6).number_format = "0.000"
    ws.cell(row=r_, column=2).number_format = "yyyy/mm/dd"
    for cc in (8, 9):
        ws.cell(row=r_, column=cc).alignment = Alignment(horizontal="center", vertical="center")
ws.conditional_formatting.add(f"I2:I{N_DEFL+1}",
    CellIsRule(operator="equal", formula=['"要確認"'], fill=WARN_FILL, font=Font(name=FONT, size=10, bold=True, color="B5326A")))
ws.conditional_formatting.add(f"I2:I{N_DEFL+1}",
    CellIsRule(operator="equal", formula=['"適合"'], fill=OK_FILL, font=Font(name=FONT, size=10, color="2C5F2D")))
ws.freeze_panes = "A2"
note(ws, f"A{N_DEFL+3}", "磁気力 F (mN) = 質量m(g) × 9.80665 × tan(θ)　／　ASTM F2052 は θ < 45°（磁気力が重力以下）を一般に許容可能とする")

# =====================================================================
# 03_偏向角_集計
# =====================================================================
ws = sheet("03_偏向角_集計")
ws["A1"] = "偏向角の集計（02_偏向角_記録 から自動計算。入力不要）"
ws["A1"].font = f(13, True, "0F1626")
ws["A2"] = "SDは標本標準偏差、95%CI半幅は t分布による。判定は同一試料・同一装置で 45° 以上の測定が1件でもあれば「要確認」。"
ws["A2"].font = f(9, False, "6B7280")
sub = [("試料ID", 11), ("製品記号", 12), ("カテゴリ", 17),
       ("n", 6), ("平均θ (°)", 10), ("SD", 8), ("95%CI半幅", 10), ("最大θ (°)", 10), ("判定", 9),
       ("n", 6), ("平均θ (°)", 10), ("SD", 8), ("95%CI半幅", 10), ("最大θ (°)", 10), ("判定", 9)]
for i, (h, w) in enumerate(sub, start=1):
    c = ws.cell(row=4, column=i, value=h)
    c.font = f(10, True, "FFFFFF"); c.fill = HDR_FILL
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = BORDER
    ws.column_dimensions[get_column_letter(i)].width = w
ws.merge_cells("D3:I3"); ws.merge_cells("J3:O3")
for rng, lab in (("D3", "1.5T"), ("J3", "3T")):
    ws[rng] = lab; ws[rng].font = f(11, True, "FFFFFF"); ws[rng].fill = PatternFill("solid", fgColor="1B9AAA")
    ws[rng].alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[4].height = 30

M = f"'02_偏向角_記録'"
DR, CR, GR = f"{M}!$D$2:$D${N_DEFL+1}", f"{M}!$C$2:$C${N_DEFL+1}", f"{M}!$G$2:$G${N_DEFL+1}"
for k in range(24):
    r = 5 + k; src = 2 + k
    ws.cell(row=r, column=1, value=f"=IF('01_試料マスタ'!A{src}=\"\",\"\",'01_試料マスタ'!A{src})")
    ws.cell(row=r, column=2, value=f"=IF('01_試料マスタ'!B{src}=\"\",\"\",'01_試料マスタ'!B{src})")
    ws.cell(row=r, column=3, value=f"=IF('01_試料マスタ'!C{src}=\"\",\"\",'01_試料マスタ'!C{src})")
    for base, dev in ((4, "1.5T"), (10, "3T")):
        L = {j: get_column_letter(base + j) for j in range(6)}
        nC, mC, sC = f"{L[0]}{r}", f"{L[1]}{r}", f"{L[2]}{r}"
        blank = f'OR($A{r}="",{nC}="",{nC}=0)'
        ws.cell(row=r, column=base + 0,
            value=f'=IF($A{r}="","",COUNTIFS({DR},$A{r},{CR},"{dev}",{GR},">=0"))')
        ws.cell(row=r, column=base + 1,
            value=f'=IFERROR(IF({blank},"",ROUND(AVERAGEIFS({GR},{DR},$A{r},{CR},"{dev}"),2)),"")').number_format = "0.00"
        ws.cell(row=r, column=base + 2,
            value=(f'=IFERROR(IF(OR($A{r}="",{nC}="",{nC}<2),"",'
                   f'ROUND(SQRT(SUMPRODUCT(({DR}=$A{r})*({CR}="{dev}")*({GR}<>"")'
                   f'*({GR}-{mC})^2)/({nC}-1)),2)),"")')).number_format = "0.00"
        ws.cell(row=r, column=base + 3,
            value=(f'=IFERROR(IF(OR($A{r}="",{nC}="",{nC}<2),"",'
                   f'ROUND(TINV(0.05,{nC}-1)*{sC}/SQRT({nC}),2)),"")')).number_format = "0.00"
        ws.cell(row=r, column=base + 4,
            value=(f'=IFERROR(IF({blank},"",'
                   f'ROUND(SUMPRODUCT(MAX(({DR}=$A{r})*({CR}="{dev}")*{GR})),2)),"")')).number_format = "0.00"
        ws.cell(row=r, column=base + 5,
            value=(f'=IFERROR(IF({blank},"",IF(COUNTIFS({DR},$A{r},{CR},"{dev}",{GR},">=45")>0,'
                   f'"要確認","適合")),"")'))
    for col in range(1, 16):
        c = ws.cell(row=r, column=col)
        c.border = BORDER; c.font = f(10)
        if col > 3:
            c.fill = CALC_FILL; c.alignment = Alignment(horizontal="center", vertical="center")
        else:
            c.fill = CALC_FILL
for colL in ("I", "O"):
    ws.conditional_formatting.add(f"{colL}5:{colL}28",
        CellIsRule(operator="equal", formula=['"要確認"'], fill=WARN_FILL, font=Font(name=FONT, size=10, bold=True, color="B5326A")))
    ws.conditional_formatting.add(f"{colL}5:{colL}28",
        CellIsRule(operator="equal", formula=['"適合"'], fill=OK_FILL, font=Font(name=FONT, size=10, color="2C5F2D")))
ws.freeze_panes = "D5"
note(ws, "A30", "試料IDは 01_試料マスタ の先頭24件を自動で参照しています。25件以上を扱う場合は最終行を下方向にコピーしてください。")

# =====================================================================
# 04_アーチファクト_記録  (案B / ASTM F2119)
# =====================================================================
ws = sheet("04_アーチファクト_記録")
cols = [("測定ID", 11, "in"), ("測定日", 12, "in"), ("装置", 9, "in"), ("シーケンス", 12, "in"),
        ("TE (ms)", 9, "in"), ("受信BW (Hz/px)", 13, "in"), ("PE方向", 10, "in"),
        ("距離 (cm)", 10, "in"), ("配置角度 (°)", 11, "in"), ("試料ID", 11, "in"), ("反復No", 8, "in"),
        ("信号消失面積 (mm2)", 15, "in"), ("最大径 (mm)", 11, "in"), ("測定者", 11, "in"), ("備考", 28, "in")]
headers(ws, cols); paint(ws, cols, 2, N_ART + 1)
example(ws, cols, 2, ["A-0001", datetime.date(2026, 2, 10), "3T", "GRE", 15, 250, "A-P", 0, 0, "S-01", 1, 121.4, 14.2, "山田", "ASTM F2119：参照画像との差30%超を閾値"])
dv(ws, f"C2:C{N_ART+1}", ["1.5T", "3T"])
dv(ws, f"D2:D{N_ART+1}", ["SE", "TSE", "GRE", "EPI", "脂肪抑制"])
dv(ws, f"G2:G{N_ART+1}", ["A-P", "R-L", "S-I"])
dv(ws, f"I2:I{N_ART+1}", ["0", "45", "90"])
dv_range(ws, f"J2:J{N_ART+1}", f"='01_試料マスタ'!$A$2:$A${N_MASTER+1}")
for r_ in range(2, N_ART + 2):
    ws.cell(row=r_, column=2).number_format = "yyyy/mm/dd"
    ws.cell(row=r_, column=12).number_format = "0.0"
    ws.cell(row=r_, column=13).number_format = "0.0"
ws.freeze_panes = "A2"
note(ws, f"A{N_ART+3}", "アーチファクトの定義（ASTM F2119）：試料なしの参照画像との画素値の差が30%を超える領域。ImageJ で閾値処理して面積・最大径を計測。")

# =====================================================================
# 05_アーチファクト_集計
# =====================================================================
ws = sheet("05_アーチファクト_集計")
ws["A1"] = "アーチファクトの集計（04_アーチファクト_記録 から自動計算。入力不要）"
ws["A1"].font = f(13, True, "0F1626")
A = "'04_アーチファクト_記録'"
dev_r, seq_r = f"{A}!$C$2:$C${N_ART+1}", f"{A}!$D$2:$D${N_ART+1}"
te_r, dis_r  = f"{A}!$E$2:$E${N_ART+1}", f"{A}!$H$2:$H${N_ART+1}"
ang_r, ar_r  = f"{A}!$I$2:$I${N_ART+1}", f"{A}!$L$2:$L${N_ART+1}"

def block(title, top, rowlabels, labelhdr, crit_builder, subtitle=None):
    ws.cell(row=top, column=1, value=title).font = f(12, True, "1B9AAA")
    if subtitle:
        ws.cell(row=top, column=4, value=subtitle).font = f(9, False, "6B7280")
    hdr = [labelhdr, "n", "平均面積 (mm2)", "SD", "n", "平均面積 (mm2)", "SD"]
    widths = [17, 7, 15, 9, 7, 15, 9]
    for i, (h, w) in enumerate(zip(hdr, widths), start=1):
        c = ws.cell(row=top + 2, column=i, value=h)
        c.font = f(10, True, "FFFFFF"); c.fill = HDR_FILL; c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if ws.column_dimensions[get_column_letter(i)].width in (None, 13):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.merge_cells(start_row=top + 1, start_column=2, end_row=top + 1, end_column=4)
    ws.merge_cells(start_row=top + 1, start_column=5, end_row=top + 1, end_column=7)
    for col, lab in ((2, "1.5T"), (5, "3T")):
        c = ws.cell(row=top + 1, column=col, value=lab)
        c.font = f(11, True, "FFFFFF"); c.fill = PatternFill("solid", fgColor="1B9AAA")
        c.alignment = Alignment(horizontal="center", vertical="center")
    for k, lab in enumerate(rowlabels):
        r = top + 3 + k
        ws.cell(row=r, column=1, value=lab).font = f(10, True)
        ws.cell(row=r, column=1).fill = CALC_FILL; ws.cell(row=r, column=1).border = BORDER
        for base, dev in ((2, "1.5T"), (5, "3T")):
            cr_cnt, cr_avg, cr_sp = crit_builder(dev, r)
            nC = f"{get_column_letter(base)}{r}"; mC = f"{get_column_letter(base+1)}{r}"
            ws.cell(row=r, column=base, value=f"=COUNTIFS({cr_cnt})")
            ws.cell(row=r, column=base + 1,
                value=f'=IFERROR(IF({nC}=0,"",ROUND(AVERAGEIFS({cr_avg}),1)),"")').number_format = "0.0"
            ws.cell(row=r, column=base + 2,
                value=f'=IFERROR(IF({nC}<2,"",ROUND(SQRT(SUMPRODUCT({cr_sp}*({ar_r}-{mC})^2)/({nC}-1)),1)),"")').number_format = "0.0"
            for j in range(3):
                c = ws.cell(row=r, column=base + j)
                c.fill = CALC_FILL; c.border = BORDER; c.font = f(10)
                c.alignment = Alignment(horizontal="center", vertical="center")
    return top + 3 + len(rowlabels)

# (a) シーケンス別（距離0cm・角度0°）
def crit_seq(dev, r):
    lab = f"$A{r}"
    cnt = f'{dev_r},"{dev}",{seq_r},{lab},{dis_r},0,{ang_r},0,{ar_r},">=0"'
    avg = f'{ar_r},{dev_r},"{dev}",{seq_r},{lab},{dis_r},0,{ang_r},0'
    sp  = f'({dev_r}="{dev}")*({seq_r}={lab})*({dis_r}=0)*({ang_r}=0)*({ar_r}<>"")'
    return cnt, avg, sp
end = block("① シーケンス別", 3, ["SE", "TSE", "GRE", "EPI", "脂肪抑制"], "シーケンス", crit_seq,
            "条件：距離 0 cm・配置角度 0°")

# (b) TE別（GRE・距離0cm）
def crit_te(dev, r):
    lab = f"$A{r}"
    cnt = f'{dev_r},"{dev}",{seq_r},"GRE",{te_r},{lab},{dis_r},0,{ar_r},">=0"'
    avg = f'{ar_r},{dev_r},"{dev}",{seq_r},"GRE",{te_r},{lab},{dis_r},0'
    sp  = f'({dev_r}="{dev}")*({seq_r}="GRE")*({te_r}={lab})*({dis_r}=0)*({ar_r}<>"")'
    return cnt, avg, sp
end = block("② TE別（GRE・距離 0 cm）", end + 2, [5, 10, 15, 20, 30], "TE (ms)", crit_te,
            "TEの水準は実際の撮影条件に合わせて書き換えてください")

# (c) 距離別（GRE）
def crit_dis(dev, r):
    lab = f"$A{r}"
    cnt = f'{dev_r},"{dev}",{seq_r},"GRE",{dis_r},{lab},{ar_r},">=0"'
    avg = f'{ar_r},{dev_r},"{dev}",{seq_r},"GRE",{dis_r},{lab}'
    sp  = f'({dev_r}="{dev}")*({seq_r}="GRE")*({dis_r}={lab})*({ar_r}<>"")'
    return cnt, avg, sp
end = block("③ 距離別（GRE）── 部位別運用の根拠になる表", end + 2, [0, 2, 5, 10, 20], "距離 (cm)", crit_dis,
            "撮影部位からの距離。ここが本研究の核心")
note(ws, f"A{end+2}", "各表の左端ラベルは入力値です。実際に撮影した水準に合わせて書き換えると、集計も追従します。", "B5326A")

# =====================================================================
# 06_発熱_記録  (案C)
# =====================================================================
ws = sheet("06_発熱_記録")
cols = [("測定ID", 11, "in"), ("測定日", 12, "in"), ("装置", 9, "in"), ("シーケンス", 12, "in"),
        ("SAR (W/kg)", 11, "in"), ("撮像時間 (min)", 12, "in"), ("試料ID", 11, "in"),
        ("40℃", 8, "in"), ("45℃", 8, "in"), ("50℃", 8, "in"), ("55℃", 8, "in"), ("60℃", 8, "in"),
        ("推定最高到達温度", 16, "calc"), ("備考", 30, "in")]
headers(ws, cols); paint(ws, cols, 2, N_HEAT + 1)
example(ws, cols, 2, ["H-001", datetime.date(2026, 2, 17), "3T", "TSE", 3.2, 30, "S-01",
                      "未変色", "未変色", "未変色", "未変色", "未変色", None, "不可逆性示温ラベル"])
for c_ in "HIJKL":
    dv(ws, f"{c_}2:{c_}{N_HEAT+1}", ["未変色", "変色"])
dv(ws, f"C2:C{N_HEAT+1}", ["1.5T", "3T"])
dv(ws, f"D2:D{N_HEAT+1}", ["SE", "TSE", "GRE", "EPI", "脂肪抑制"])
dv_range(ws, f"G2:G{N_HEAT+1}", f"='01_試料マスタ'!$A$2:$A${N_MASTER+1}")
for r_ in range(2, N_HEAT + 2):
    ws.cell(row=r_, column=13, value=(
        f'=IF(H{r_}="","",IF(L{r_}="変色","60℃以上",IF(K{r_}="変色","55〜60℃",'
        f'IF(J{r_}="変色","50〜55℃",IF(I{r_}="変色","45〜50℃",IF(H{r_}="変色","40〜45℃","40℃未満"))))))'))
    ws.cell(row=r_, column=13).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=r_, column=2).number_format = "yyyy/mm/dd"
ws.conditional_formatting.add(f"M2:M{N_HEAT+1}",
    CellIsRule(operator="notEqual", formula=['"40℃未満"'], fill=WARN_FILL, font=Font(name=FONT, size=10, bold=True, color="B5326A")))
ws.freeze_panes = "A2"
note(ws, f"A{N_HEAT+3}", "不可逆性示温ラベルは冷えても色が戻らないため、ガントリーから出すまでの遅延に影響されず最高到達温度が残る。")
note(ws, f"A{N_HEAT+4}", "限界：温度の時間変化は取得できない。ファントムは血流による冷却を再現していないため、生体より高温側に出る（安全側の評価）。")

# =====================================================================
# 07_延期症例  (案E-1)
# =====================================================================
ws = sheet("07_延期症例")
cols = [("症例No", 11, "in"), ("依頼日", 12, "in"), ("初回予約日", 12, "in"), ("実施日", 12, "in"),
        ("撮影部位", 13, "in"), ("ネイル種別", 14, "in"), ("延期", 8, "in"), ("延期理由", 22, "in"),
        ("未受検", 10, "in"), ("依頼→実施 日数", 14, "calc"), ("延期例の日数", 13, "calc"),
        ("非延期例の日数", 13, "calc"), ("備考", 24, "in")]
headers(ws, cols); paint(ws, cols, 2, N_CASE + 1)
example(ws, cols, 2, ["C-001", datetime.date(2026, 1, 6), datetime.date(2026, 1, 13), datetime.date(2026, 2, 4), "頭部", "マグネットネイル",
                      "有", "ネイル除去のため", "非該当", None, None, None, ""])
dv(ws, f"E2:E{N_CASE+1}", ["頭部", "体幹部", "手・手関節", "四肢（手以外）", "その他"])
dv(ws, f"F2:F{N_CASE+1}", ["マグネットネイル", "ジェルネイル", "マニキュア", "不明", "なし"])
dv(ws, f"G2:G{N_CASE+1}", ["有", "無"])
dv(ws, f"I2:I{N_CASE+1}", ["該当", "非該当"])
for r_ in range(2, N_CASE + 2):
    ws.cell(row=r_, column=10, value=f'=IF(OR(B{r_}="",D{r_}=""),"",D{r_}-B{r_})').number_format = "0"
    ws.cell(row=r_, column=11, value=f'=IF(AND(G{r_}="有",J{r_}<>""),J{r_},"")').number_format = "0"
    ws.cell(row=r_, column=12, value=f'=IF(AND(G{r_}="無",J{r_}<>""),J{r_},"")').number_format = "0"
    for cc in (2, 3, 4):
        ws.cell(row=r_, column=cc).number_format = "yyyy/mm/dd"
    for cc in (10, 11, 12):
        ws.cell(row=r_, column=cc).alignment = Alignment(horizontal="center", vertical="center")
ws.freeze_panes = "A2"
note(ws, f"A{N_CASE+3}", "「未受検」は、延期後に再予約されず検査が実施されないまま終わった症例。最も重い指標なので必ず拾ってください。", "B5326A")

# =====================================================================
# 08_延期集計
# =====================================================================
ws = sheet("08_延期集計")
ws["A1"] = "延期の実態（07_延期症例 から自動計算。入力不要）"
ws["A1"].font = f(13, True, "0F1626")
ws["A2"] = "発表の導入スライドはこの表の数値を使います。"
ws["A2"].font = f(9, False, "6B7280")
for col, w in (("A", 30), ("B", 14), ("C", 46)):
    ws.column_dimensions[col].width = w
C = "'07_延期症例'"
g_r = f"{C}!$G$2:$G${N_CASE+1}"; a_r = f"{C}!$A$2:$A${N_CASE+1}"
e_r = f"{C}!$E$2:$E${N_CASE+1}"; i_r = f"{C}!$I$2:$I${N_CASE+1}"
j_r = f"{C}!$J$2:$J${N_CASE+1}"; k_r = f"{C}!$K$2:$K${N_CASE+1}"; l_r = f"{C}!$L$2:$L${N_CASE+1}"
items = [
    ("対象症例数", f"=COUNTA({a_r})", "0", "07に記録した全症例"),
    ("延期件数", f'=COUNTIF({g_r},"有")', "0", "ネイルを理由に延期した件数"),
    ("延期率", f'=IFERROR(B5/B4,"")', "0.0%", "延期件数 ÷ 対象症例数"),
    ("未受検（脱落）件数", f'=COUNTIF({i_r},"該当")', "0", "延期後に再予約されず終わった症例。最重要指標"),
    ("依頼→実施 日数　中央値（全体）", f"=IFERROR(MEDIAN({j_r}),\"\")", "0.0", "全症例"),
    ("　同　中央値（延期例）", f"=IFERROR(MEDIAN({k_r}),\"\")", "0.0", "延期した症例のみ。発表の主役の数字"),
    ("　同　中央値（非延期例）", f"=IFERROR(MEDIAN({l_r}),\"\")", "0.0", "延期しなかった症例のみ"),
    ("延期による遅れ（差）", '=IFERROR(IF(OR(B9="",B10=""),"",B9-B10),"")', "0.0", "延期例 − 非延期例。これが「診断の遅れ」"),
]
r = 4
for lab, fml, nf, memo in items:
    ws.cell(row=r, column=1, value=lab).font = f(10, True)
    ws.cell(row=r, column=1).fill = CALC_FILL; ws.cell(row=r, column=1).border = BORDER
    c = ws.cell(row=r, column=2, value=fml); c.number_format = nf
    c.fill = CALC_FILL; c.border = BORDER; c.font = f(11, True, "B5326A")
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=r, column=3, value=memo).font = f(9, False, "6B7280")
    r += 1

ws.cell(row=13, column=1, value="部位別の内訳").font = f(12, True, "1B9AAA")
ws.cell(row=13, column=3, value="案Bの距離依存性の結果と突き合わせ、「回避できたはずの遅れ」を出します").font = f(9, False, "6B7280")
for i, h in enumerate(["撮影部位", "延期件数", "構成比"], start=1):
    c = ws.cell(row=14, column=i, value=h)
    c.font = f(10, True, "FFFFFF"); c.fill = HDR_FILL; c.border = BORDER
    c.alignment = Alignment(horizontal="center", vertical="center")
for k, part in enumerate(["頭部", "体幹部", "手・手関節", "四肢（手以外）", "その他"]):
    r = 15 + k
    ws.cell(row=r, column=1, value=part).font = f(10, True)
    ws.cell(row=r, column=2, value=f'=COUNTIFS({e_r},$A{r},{g_r},"有")').number_format = "0"
    ws.cell(row=r, column=3, value=f'=IFERROR(B{r}/$B$5,"")').number_format = "0.0%"
    for cc in (1, 2, 3):
        cell = ws.cell(row=r, column=cc); cell.fill = CALC_FILL; cell.border = BORDER; cell.font = f(10)
        if cc > 1:
            cell.alignment = Alignment(horizontal="center", vertical="center")
ws.cell(row=20, column=1, value="合計").font = f(10, True)
ws.cell(row=20, column=2, value="=SUM(B15:B19)").number_format = "0"
ws.cell(row=20, column=3, value='=IFERROR(B20/$B$5,"")').number_format = "0.0%"
for cc in (1, 2, 3):
    cell = ws.cell(row=20, column=cc); cell.fill = SUB_FILL; cell.border = BORDER; cell.font = f(10, True)
note(ws, "A22", "頭部・体幹部の延期件数は、案Bで「画像への影響は検出されなかった」と示せた場合、そのまま「回避できたはずの遅れ」として提示できます。", "B5326A")

# =====================================================================
# 09_スクリーニング検証  (案E-3)
# =====================================================================
ws = sheet("09_スクリーニング検証")
ws["A1"] = "簡易磁石テストの性能評価（案E-3）"
ws["A1"].font = f(13, True, "0F1626")
ws["A2"] = "「基準判定」は案A・案Bの結果から決めた正解ラベル。磁石テストの結果と突き合わせます。"
ws["A2"].font = f(9, False, "6B7280")
cols = [("試料ID", 11, "in"), ("磁石テスト結果", 14, "in"), ("基準判定", 14, "in"), ("判定者", 11, "in"), ("備考", 30, "in")]
headers(ws, cols, row=4); paint(ws, cols, 5, N_SCR + 4)
example(ws, cols, 5, ["S-01", "陽性", "磁性あり", "山田", "冷蔵庫マグネット・5mm"])
dv(ws, f"B5:B{N_SCR+4}", ["陽性", "陰性"])
dv(ws, f"C5:C{N_SCR+4}", ["磁性あり", "磁性なし"])
dv_range(ws, f"A5:A{N_SCR+4}", f"='01_試料マスタ'!$A$2:$A${N_MASTER+1}")
ws.freeze_panes = "A5"

br, cr = f"$B$5:$B${N_SCR+4}", f"$C$5:$C${N_SCR+4}"
ws.cell(row=4, column=7, value="2×2 表").font = f(12, True, "1B9AAA")
grid = [("", "磁性あり", "磁性なし"),
        ("テスト陽性", f'=COUNTIFS({br},"陽性",{cr},"磁性あり")', f'=COUNTIFS({br},"陽性",{cr},"磁性なし")'),
        ("テスト陰性", f'=COUNTIFS({br},"陰性",{cr},"磁性あり")', f'=COUNTIFS({br},"陰性",{cr},"磁性なし")')]
for i, row_ in enumerate(grid):
    for j, v in enumerate(row_):
        c = ws.cell(row=5 + i, column=7 + j, value=v)
        c.border = BORDER; c.alignment = Alignment(horizontal="center", vertical="center")
        if i == 0 or j == 0:
            c.font = f(10, True, "FFFFFF"); c.fill = HDR_FILL
        else:
            c.font = f(11, True); c.fill = CALC_FILL
for col in ("G", "H", "I"):
    ws.column_dimensions[col].width = 13
ws.cell(row=5, column=7).fill = PatternFill("solid", fgColor="FFFFFF")

ws.cell(row=9, column=7, value="指標").font = f(12, True, "1B9AAA")
metrics = [("感度（見逃しの少なさ）", '=IFERROR(H6/(H6+H7),"")', "磁性ありを正しく陽性と判定できた割合"),
           ("特異度（空振りの少なさ）", '=IFERROR(I7/(I7+I6),"")', "磁性なしを正しく陰性と判定できた割合"),
           ("陽性的中率（PPV）", '=IFERROR(H6/(H6+I6),"")', "陽性と出たもののうち実際に磁性ありの割合"),
           ("陰性的中率（NPV）", '=IFERROR(I7/(I7+H7),"")', "陰性と出たもののうち実際に磁性なしの割合"),
           ("正確度", '=IFERROR((H6+I7)/(H6+I6+H7+I7),"")', "全体の一致率")]
for k, (lab, fml, memo) in enumerate(metrics):
    r = 10 + k
    ws.cell(row=r, column=7, value=lab).font = f(10, True)
    ws.cell(row=r, column=7).fill = CALC_FILL; ws.cell(row=r, column=7).border = BORDER
    c = ws.cell(row=r, column=8, value=fml); c.number_format = "0.0%"
    c.font = f(11, True, "B5326A"); c.fill = CALC_FILL; c.border = BORDER
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=r, column=9, value=memo).font = f(9, False, "6B7280")
ws.column_dimensions["I"].width = 42
note(ws, "G16", "受付で使うテストなので、重視すべきは感度（磁性ありを見逃さないこと）。", "B5326A")

# =====================================================================
# 10_撮影ログ
# =====================================================================
ws = sheet("10_撮影ログ")
cols = [("実施日", 12, "in"), ("装置", 9, "in"), ("開始", 9, "in"), ("終了", 9, "in"),
        ("所要 (h)", 9, "calc"), ("目的・実施内容", 34, "in"), ("立会者", 20, "in"),
        ("試料点数", 9, "in"), ("特記事項・中断の有無", 34, "in")]
headers(ws, cols); paint(ws, cols, 2, N_LOG + 1)
example(ws, cols, 2, [datetime.date(2026, 2, 10), "3T", datetime.time(18, 30), datetime.time(21, 0), None, "案B 距離依存性 撮影", "山田・佐藤", 12, "異常なし"])
dv(ws, f"B2:B{N_LOG+1}", ["1.5T", "3T"])
for r_ in range(2, N_LOG + 2):
    ws.cell(row=r_, column=5, value=f'=IF(OR(C{r_}="",D{r_}=""),"",ROUND((D{r_}-C{r_})*24,1))').number_format = "0.0"
    ws.cell(row=r_, column=1).number_format = "yyyy/mm/dd"
    for cc in (3, 4):
        ws.cell(row=r_, column=cc).number_format = "hh:mm"
    ws.cell(row=r_, column=5).alignment = Alignment(horizontal="center", vertical="center")
ws.freeze_panes = "A2"
note(ws, f"A{N_LOG+3}", "立会者は必ず2名以上。時間外の単独作業は行わないこと。")

# =====================================================================
# 11_安全チェックリスト
# =====================================================================
ws = sheet("11_安全チェックリスト")
ws["A1"] = "安全チェックリスト（撮影のたびに確認）"
ws["A1"].font = f(13, True, "0F1626")
ws["A2"] = "実施日と実施者を記入し、各項目を確認してから検査室に入ってください。このシートを印刷して使うこともできます。"
ws["A2"].font = f(9, False, "6B7280")
ws["A4"] = "実施日"; ws["A4"].font = f(10, True)
ws["B4"] = ""; ws["B4"].fill = IN_FILL; ws["B4"].border = BORDER
ws["C4"] = "実施者"; ws["C4"].font = f(10, True)
ws["D4"] = ""; ws["D4"].fill = IN_FILL; ws["D4"].border = BORDER
for col, w in (("A", 12), ("B", 16), ("C", 10), ("D", 20), ("E", 78), ("F", 8)):
    ws.column_dimensions[col].width = w
checks = [
    ("実施前", "強力磁石（ネオジム磁石）を検査室に持ち込んでいないこと。模様付けと硬化は検査室外で完了させる"),
    ("実施前", "金属を含む温度計・工具・治具を持ち込んでいないこと。治具は樹脂・木材など非磁性材であること"),
    ("実施前", "全試料に回収用の索を付け、確実に固定したこと"),
    ("実施前", "立会者が2名以上そろっていること"),
    ("実施前", "中止基準を共有したこと（異音／試料のズレ／示温ラベルの想定外の変色 → 即中断）"),
    ("実施前", "持ち込み物を一覧で確認したこと"),
    ("実施中", "人体（自分の手を含む）の撮影を行っていないこと。ファントムのみであること"),
    ("実施中", "試料の位置ずれがないことを撮影ごとに確認したこと"),
    ("実施後", "全試料を回収し、点数が一致することを確認したこと"),
    ("実施後", "撮影条件と実施内容を 10_撮影ログ に記録したこと"),
    ("実施後", "検査室内に持ち込み物が残っていないことを確認したこと"),
]
for i, h in enumerate(["区分", "", "", "", "確認項目", "確認"], start=1):
    if h:
        c = ws.cell(row=6, column=i, value=h)
        c.font = f(10, True, "FFFFFF"); c.fill = HDR_FILL; c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center")
ws.cell(row=6, column=1, value="区分").font = f(10, True, "FFFFFF")
for k, (cat, txt) in enumerate(checks):
    r = 7 + k
    c1 = ws.cell(row=r, column=1, value=cat); c1.font = f(10, True); c1.fill = CALC_FILL
    c1.alignment = Alignment(horizontal="center", vertical="center"); c1.border = BORDER
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    c2 = ws.cell(row=r, column=2, value=txt); c2.font = f(10)
    c2.alignment = Alignment(vertical="center", wrap_text=True); c2.border = BORDER
    for cc in range(3, 6):
        ws.cell(row=r, column=cc).border = BORDER
    c3 = ws.cell(row=r, column=6); c3.fill = IN_FILL; c3.border = BORDER
    c3.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[r].height = 30
dv(ws, f"F7:F{6+len(checks)}", ["✓", "該当なし"])
note(ws, f"A{8+len(checks)}", "人体でのMRI撮影は行わない。ファントムのみ。これは例外なく守ってください。", "B5326A", 10)

out = sys.argv[1] if len(sys.argv) > 1 else "magnet-nail-mri_workbook.xlsx"
wb.save(out)
print("saved:", out, "| sheets:", len(wb.sheetnames))
